#!/usr/bin/env python3
"""
VC Breakout v4.4 — Excel Builder
=================================
Reads scored.json, produces a formatted .xlsx with 4 tabs:
  - One per AE (flat ranked list, 18 columns)
  - Weekly Movers (10 columns)

Usage:
    python3 build_xlsx.py --input scored.json --output vc_breakout_YYYY-MM-DD.xlsx
"""

import json
import argparse
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
except ImportError:
    import sys
    print("openpyxl required: pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)


HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=10)
HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
HEADER_ALIGN = Alignment(horizontal="center", vertical="center", wrap_text=True)

SCORE_GREEN = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
SCORE_YELLOW = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
SCORE_RED = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

THIN_BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

AE_COLUMNS = [
    "Rank", "Δ Rank", "Account Name", "Decision Country", "Startup Type",
    "Funding Stage", "Breakout Score", "D1 PIV", "D2 Funding", "D3 External",
    "D4 VC Tier", "D5 Founder", "PIV 90d €", "PIV Growth %", "AI Flag",
    "Months Since Round", "Key Signal", "Recommended Action",
]

MOVERS_COLUMNS = [
    "Account Name", "AE", "Decision Country", "Startup Type", "Trigger",
    "Old Rank", "New Rank", "Δ", "Score Δ", "Recommended Action",
]

AE_TAB_NAMES = {
    "Jeanne": "🇫🇷 Jeanne",
    "Fionn": "🇬🇧 Fionn",
    "Rami": "🇩🇪 Rami",
}


def format_delta(delta):
    if delta == "NEW":
        return "NEW"
    if isinstance(delta, (int, float)):
        if delta > 0:
            return f"+{delta}"
        return str(delta)
    return str(delta)


def build_ae_row(acc: dict) -> list:
    ds = acc.get("dimension_scores", {})
    return [
        acc.get("rank"),
        format_delta(acc.get("delta_rank", "NEW")),
        acc.get("account_name", "—"),
        acc.get("country", "—"),
        acc.get("startup_type", "—"),
        acc.get("funding_stage", "—"),
        acc.get("breakout_score", 0),
        round(ds.get("piv_growth", 0), 1),
        round(ds.get("funding_momentum", 0), 1),
        round(ds.get("ext_momentum", 0), 1),
        round(ds.get("vc_tier", 0), 1),
        round(ds.get("founder_pedigree", 0), 1),
        acc.get("piv_last_90d_eur"),
        acc.get("piv_90d_growth_pct"),
        "AI" if acc.get("ai_flag") else "",
        acc.get("months_since_round", "—"),
        acc.get("key_signal", "—"),
        acc.get("recommended_action", "—"),
    ]


def build_movers_row(m: dict) -> list:
    return [
        m.get("account_name", "—"),
        m.get("ae", "—"),
        m.get("country", "—"),
        m.get("startup_type", "—"),
        m.get("trigger", "—"),
        m.get("old_rank"),
        m.get("new_rank"),
        m.get("delta"),
        m.get("score_delta"),
        m.get("recommended_action", "—"),
    ]


def style_header(ws, columns):
    for col_idx, col_name in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=col_name)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
        cell.border = THIN_BORDER


def apply_score_fill(ws, score_col_idx, start_row, end_row):
    for row in range(start_row, end_row + 1):
        cell = ws.cell(row=row, column=score_col_idx)
        val = cell.value
        if val is None:
            continue
        if val >= 70:
            cell.fill = SCORE_GREEN
        elif val >= 40:
            cell.fill = SCORE_YELLOW
        else:
            cell.fill = SCORE_RED


def auto_width(ws, columns):
    for col_idx, col_name in enumerate(columns, 1):
        max_len = len(col_name)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_len + 2, 40)


def build_workbook(scored_data: dict, output_path: str):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    accounts = scored_data.get("scored_accounts", [])
    movers = scored_data.get("weekly_movers", [])

    ae_order = ["Jeanne", "Fionn", "Rami"]
    for ae in ae_order:
        ae_accounts = [a for a in accounts if a["ae"] == ae]
        ae_accounts.sort(key=lambda a: a.get("rank", 9999))

        tab_name = AE_TAB_NAMES.get(ae, ae)
        ws = wb.create_sheet(title=tab_name)
        style_header(ws, AE_COLUMNS)

        for row_idx, acc in enumerate(ae_accounts, 2):
            row_data = build_ae_row(acc)
            for col_idx, val in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=val)

        score_col = AE_COLUMNS.index("Breakout Score") + 1
        apply_score_fill(ws, score_col, 2, len(ae_accounts) + 1)
        auto_width(ws, AE_COLUMNS)
        ws.freeze_panes = "A2"

    ws_movers = wb.create_sheet(title="⚡ Weekly Movers")
    style_header(ws_movers, MOVERS_COLUMNS)
    for row_idx, m in enumerate(movers, 2):
        row_data = build_movers_row(m)
        for col_idx, val in enumerate(row_data, 1):
            ws_movers.cell(row=row_idx, column=col_idx, value=val)
    auto_width(ws_movers, MOVERS_COLUMNS)
    ws_movers.freeze_panes = "A2"

    wb.save(output_path)
    print(f"Written: {output_path} ({len(accounts)} accounts, {len(movers)} movers)")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--output", required=True)
    args = parser.parse_args()

    with open(args.input) as f:
        data = json.load(f)

    build_workbook(data, args.output)


if __name__ == "__main__":
    main()
